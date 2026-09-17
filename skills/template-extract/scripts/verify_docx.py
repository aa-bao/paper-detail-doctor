#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""verify_docx.py —— 回读生成的 .docx，逐项校验样式是否与 Spec 一致。

这是“回读校验”环节：不信任写入过程，直接把成品打开读数，
对比 spec 期望值，输出差异清单。任何 mismatch 都值得人工看一眼。
"""
import argparse
import json
import os
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C  # noqa: E402


def _pt(v):
    try:
        return round(v.pt, 2)
    except Exception:
        return None


def read_style(doc, name):
    try:
        st = doc.styles[name]
    except KeyError:
        return None
    out = {"exists": True}
    rpr = st.element.find(qn("w:rPr"))
    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            out["cn_font"] = rf.get(qn("w:eastAsia"))
            out["en_font"] = rf.get(qn("w:ascii"))
    out["size_pt"] = _pt(st.font.size)
    out["bold"] = st.font.bold
    pf = st.paragraph_format
    out["align"] = pf.alignment.name.lower() if pf.alignment is not None else None
    out["space_before_pt"] = _pt(pf.space_before)
    out["space_after_pt"] = _pt(pf.space_after)
    ls = pf.line_spacing
    try:
        out["line_spacing"] = round(float(ls), 3) if ls is not None else None
    except Exception:
        out["line_spacing"] = _pt(ls)
    pPr = st.element.find(qn("w:pPr"))
    if pPr is not None:
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            flc = ind.get(qn("w:firstLineChars"))
            lc = ind.get(qn("w:leftChars"))
            out["first_line_indent_chars"] = int(flc) / 100 if flc else 0
            out["left_indent_chars"] = int(lc) / 100 if lc else 0
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            numId = numPr.find(qn("w:numId"))
            ilvl = numPr.find(qn("w:ilvl"))
            out["numbering"] = {
                "numId": numId.get(qn("w:val")) if numId is not None else None,
                "ilvl": ilvl.get(qn("w:val")) if ilvl is not None else None,
            }
    return out


SPEC_KEYS = {
    "print": ["cn_font", "en_font", "size_pt", "bold", "align",
              "space_before_pt", "space_after_pt",
              "first_line_indent_chars", "left_indent_chars"],
    "approx": {"line_spacing": 0.06},
}


def compare(expected, actual):
    diffs = []
    if actual is None:
        return [{"field": "(style)", "expected": "exists", "actual": "MISSING"}]
    for f in SPEC_KEYS["print"]:
        if f not in expected:
            continue
        e, a = expected.get(f), actual.get(f)
        # 缩进为 0 时不会写入 w:ind，读回是 None —— 语义等价于 0
        if f in ("first_line_indent_chars", "left_indent_chars"):
            if e is None:
                e = 0
            if a is None:
                a = 0
        if f == "align":
            e = C.resolve_align(e).lower() if e is not None else None
            a = (a or "").lower() or None
        if e is None and a is None:
            continue
        if isinstance(e, (int, float)) and isinstance(a, (int, float)):
            if abs(float(e) - float(a)) > 0.01:
                diffs.append({"field": f, "expected": e, "actual": a})
        elif e != a:
            if not (e in (None, "") and a in (None, "", False)):
                diffs.append({"field": f, "expected": e, "actual": a})
    e_ls = expected.get("line_spacing")
    a_ls = actual.get("line_spacing")
    if e_ls is not None and a_ls is not None:
        try:
            if abs(float(e_ls) - float(a_ls)) > 0.06:
                diffs.append({"field": "line_spacing", "expected": e_ls, "actual": a_ls})
        except Exception:
            pass
    return diffs


def style_usage(doc, spec_names=()):
    """样式体检：定义了哪些样式、被引用的哪些、哪些是死样式。

    对应 WPS/Word「样式」面板里那一堆看着碍眼的项 —— 生成的模板不允许留死样式。
    三类要分开看：
      * 被引用     —— 文档里真的在用
      * 预留       —— Spec 里定义了、留给写作时用（如四级标题、目录一/二/三级），
                      不算死样式
      * 未使用     —— 既没被引用、Spec 里也没有 —— 这才是该清理的"多余样式"
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    roots = [doc.element.body]
    for rel in list(doc.part.rels.values()):
        if getattr(rel, "is_external", False):
            continue
        if rel.reltype in (RT.HEADER, RT.FOOTER):
            try:
                roots.append(rel.target_part.element)
            except Exception:
                pass

    used = set()
    for root in roots:
        for el in root.iter():
            if el.tag in (qn("w:pStyle"), qn("w:rStyle"), qn("w:tblStyle")):
                v = el.get(qn("w:val"))
                if v:
                    used.add(v)

    defined, required = {}, set()
    for st in doc.styles:
        try:
            sid = st.style_id
        except Exception:
            continue
        if sid is None:
            continue
        defined[sid] = (st.name, str(getattr(st.type, "name", st.type)))
        if st.element.get(qn("w:default")) == "1":
            required.add(sid)

    want = {str(n).strip() for n in spec_names if n}
    reserved, unused = [], []
    for sid, (name, _t) in defined.items():
        if sid in used or sid in required:
            continue
        (reserved if name in want else unused).append(name)
    return {
        "styles_defined": len(defined),
        "styles_referenced": len(used & set(defined)),
        "styles_required_defaults": len(required),
        "styles_reserved": sorted(reserved),
        "styles_unused": sorted(unused),
        "clean": not unused,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-d", "--docx", required=True)
    ap.add_argument("-s", "--spec", required=True)
    args = ap.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)
    doc = Document(args.docx)

    report = {"docx": os.path.abspath(args.docx), "styles": {}, "page": {}, "sections": []}

    # ---- 页面
    sec0 = doc.sections[0]
    page = spec.get("page", {})
    for spec_key, attr, scale in (
        ("width_cm", "page_width", 2.54), ("height_cm", "page_height", 2.54),
        ("margin_top_cm", "top_margin", 2.54), ("margin_bottom_cm", "bottom_margin", 2.54),
        ("margin_left_cm", "left_margin", 2.54), ("margin_right_cm", "right_margin", 2.54),
    ):
        if spec_key not in page:
            continue
        exp = float(page[spec_key])
        got = float(getattr(sec0, attr).cm)
        if abs(exp - got) > 0.02:
            report["page"][spec_key] = {"expected": exp, "actual": round(got, 2)}
    page_diffs = len(report["page"])
    report["page"]["_sections"] = len(doc.sections)
    report["page"]["_checked"] = "ok" if page_diffs == 0 else "diff"

    # ---- 样式
    total_diff = 0
    for key, s in spec.get("styles", {}).items():
        name = s["display_name"]
        actual = read_style(doc, name)
        diffs = compare(s, actual)
        report["styles"][key] = {
            "display_name": name,
            "status": "OK" if not diffs else "DIFF",
            "diffs": diffs,
        }
        total_diff += len(diffs)

    # ---- 节页码格式
    for i, sec in enumerate(doc.sections):
        sectPr = sec._sectPr
        pgNumType = sectPr.find(qn("w:pgNumType"))
        report["sections"].append({
            "index": i,
            "page_number_fmt": pgNumType.get(qn("w:fmt")) if pgNumType is not None else None,
            "page_number_start": pgNumType.get(qn("w:start")) if pgNumType is not None else None,
            "has_header": bool("".join(p.text for p in sec.header.paragraphs).strip()),
            "header_text": "".join(p.text for p in sec.header.paragraphs).strip(),
        })

    report["summary"] = {
        "styles_total": len(spec.get("styles", {})),
        "styles_with_diff": sum(1 for v in report["styles"].values() if v["status"] == "DIFF"),
        "total_field_diffs": total_diff,
    }
    report["cleanliness"] = style_usage(
        doc, spec_names=[s.get("display_name") for s in spec.get("styles", {}).values()])
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
