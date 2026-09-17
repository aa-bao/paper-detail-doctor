#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""
tests/selftest_format_audit.py —— format-audit（L1）的**注入式**验收

为什么不能只跑一遍干净文档就收工
    "在一份本来就干净的稿子上跑出 0 条问题"完全不能说明检查器是好的 ——
    一个永远返回空列表的函数也能做到。所以这里反过来做：
    **人为造出每一种缺陷，确认检查器抓得到**。

六个场景（都在临时副本上，不动你的原稿）
    0 基线          干净稿上只应剩真实问题（页眉横线各节不一致），且不得有 error；
                   确认只读（哈希不变）、问题 ID 稳定
    1 page-setup    把某节左边距改成 6cm → 因缺标准**不应**新增 page-setup 问题（行数不变）
    2 page-numbering 删掉第 1 节的 w:pgNumType → 检出 page-numbering
    3 header        给某节页眉加一句不一致文字 → 检出 header（文字不一致）
    4 fonts/style   给某正文段设显式 w:sz（9pt，与样式不同）→ 检出 fonts 或 style-compliance
    5 toc           清空目录域的 instrText → 检出 toc（域失效）

怎么跑
    cd F:\Coding\Project\paper-detail-doctor
    "C:/Users/Tian/.workbuddy/binaries/python/envs/default/Scripts/python.exe" tests/selftest_format_audit.py
"""
import importlib.util
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
sys.path.insert(0, PKG)

from shared.lib.docx_ops import Doc, W                       # noqa: E402
from shared.lib.docx_scan import norm                        # noqa: E402
from shared.lib.report import validate_issue                 # noqa: E402
from shared.lib.docx_ops import ACTIONS                      # noqa: E402

DEFAULT_SAMPLE = r'F:\Coding\work\0813-城市家庭\初稿.docx'
TMP = tempfile.mkdtemp(prefix='pdd-fmt-selftest-')

_ok, _fail, _notes = 0, 0, []


def check(name, cond, detail=''):
    global _ok, _fail
    if cond:
        _ok += 1
        print('  [ok]   %s' % name)
    else:
        _fail += 1
        print('  [FAIL] %s   %s' % (name, detail))
        _notes.append('%s :: %s' % (name, detail))


def section(t):
    print('\n=== %s ===' % t)


def load_audit_mod():
    p = os.path.join(PKG, 'skills', 'format-audit', 'scripts', 'audit.py')
    spec = importlib.util.spec_from_file_location('fmt_audit', p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


AUD = None


def audit(docx, spec=None):
    return AUD.audit(docx, spec, verbose=False)['issues']


def rules_of(issues):
    return {i['rule'] for i in issues}


def fresh(src, tag):
    p = os.path.join(TMP, tag + '.docx')
    shutil.copy2(src, p)
    return p


# ---------------------------------------------------------------- 注入器

def inject_margin_6cm(docx):
    """把第 1 节左边距改成 6cm（twips: 6cm ≈ 3402 twips）。"""
    from docx.oxml.ns import qn
    d = Doc(docx)
    sec = d.doc.sections[0]
    sectPr = sec._sectPr
    pgmar = sectPr.find(W + 'pgMar')
    if pgmar is None:
        pgmar = sectPr.makeelement(W + 'pgMar', {})
        sectPr.append(pgmar)
    pgmar.set(W + 'left', '3402')          # 6cm
    d.save_atomic()
    return 'left=3402twips(6cm)'


def inject_delete_pgNumType(docx, sec_index=1):
    """删掉第 sec_index 节的 w:pgNumType。"""
    d = Doc(docx)
    sec = d.doc.sections[sec_index - 1]
    sectPr = sec._sectPr
    pg = sectPr.find(W + 'pgNumType')
    if pg is not None:
        sectPr.remove(pg)
    d.save_atomic()
    return 'sec%d pgNumType removed' % sec_index


def inject_header_inconsistent(docx, sec_index=1):
    """给第 sec_index 节页眉写一句与第 2 节不一致的文字。"""
    d = Doc(docx)
    sec = d.doc.sections[sec_index - 1]
    sec.header.is_linked_to_previous = False
    paras = sec.header.paragraphs
    if paras and (paras[0].text or '').strip():
        paras[0].text = '封面题目占位文字-XX'
    else:
        sec.header.add_paragraph('封面题目占位文字-XX')
    d.save_atomic()
    return 'sec%d header text set' % sec_index


def inject_run_sz(docx, sz_half=18, style='正文文本1'):
    """给一个「正文文本1」段落的首个 run 设显式 w:sz（9pt），与样式（12pt）不同。"""
    d = Doc(docx)
    target = None
    for p in d.doc.paragraphs:
        if p.style and p.style.name == style and p.runs:
            target = p
            break
    if target is None:
        raise RuntimeError('找不到 %s 段落' % style)
    r = target.runs[0]
    rpr = r._r.find(W + 'rPr')
    if rpr is None:
        rpr = r._r.makeelement(W + 'rPr', {})
        r._r.insert(0, rpr)
    sz = rpr.makeelement(W + 'sz', {})
    sz.set(W + 'val', str(sz_half))
    rpr.append(sz)
    d.save_atomic()
    return 'para%d run0 w:sz=%d' % (target._index if hasattr(target, '_index') else -1, sz_half)


def inject_clear_toc(docx):
    """清空目录域的 TOC instrText（使其不再是合法 TOC 域）。"""
    d = Doc(docx)
    h, region = AUD._toc_region(d.doc)
    if h is None:
        raise RuntimeError('找不到目录')
    hit = 0
    for pi in region:
        for it in d.doc.paragraphs[pi - 1]._element.iter(W + 'instrText'):
            if it.text and 'TOC' in (it.text or '').upper():
                it.text = ''
                hit += 1
    if hit == 0:
        raise RuntimeError('目录区未找到 TOC instrText')
    d.save_atomic()
    return 'cleared %d TOC instrText' % hit


# ---------------------------------------------------------------- 场景

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SAMPLE
    global AUD
    AUD = load_audit_mod()
    if not os.path.exists(src):
        print('样本不存在：%s' % src)
        return 2
    print('样本：%s\n临时目录：%s' % (src, TMP))

    # ---------------------------------------------------- 0 基线
    section('0 基线：干净稿上不该冒出 error')
    base = fresh(src, 'base')
    iss = audit(base)
    base_rules = rules_of(iss)
    base_count = len(iss)
    check('基线无 error', not [i for i in iss if i['severity'] == 'error'],
          'errors=%s' % [(i['rule'], i['actual'][:40]) for i in iss if i['severity'] == 'error'])
    check('基线只剩已知的真实问题（header 横线不一致）',
          base_rules <= {'header'}, 'rules=%s' % base_rules)
    check('每条 Issue 都满足 schema',
          all(not validate_issue(i) for i in iss),
          'bad=%s' % [validate_issue(i) for i in iss if validate_issue(i)][:2])
    check('action 名字都在 ACTIONS 里（或为空）',
          all((i.get('action') is None) or (i['action'].get('name') in ACTIONS)
              for i in iss),
          'bad=%s' % [i.get('action') for i in iss if i.get('action') and i['action'].get('name') not in ACTIONS])
    h0 = AUD._sha256(base) if False else __import__('shared.lib.ooxml_guard', fromlist=['sha256_file']).sha256_file(base)
    iss2 = audit(base)
    check('审计是只读的（哈希不变）',
          __import__('shared.lib.ooxml_guard', fromlist=['sha256_file']).sha256_file(base) == h0)
    check('问题 ID 稳定（重跑同 ID）',
          [i['id'] for i in iss] == [i['id'] for i in iss2],
          '%s vs %s' % ([i['id'] for i in iss][:2], [i['id'] for i in iss2][:2]))
    print('  [info] 基线问题条数=%d，规则分布=%s' % (base_count, base_rules))

    # ---------------------------------------------------- 1 page-setup（无标准→不判）
    section('1 注入：第 1 节左边距改 6cm → 因缺标准不应新增 page-setup')
    f = fresh(src, 'pagesetup')
    info = inject_margin_6cm(f)
    iss = audit(f)
    check('注入后无 page-setup 问题（page_margins_cm 为 None 不判）',
          'page-setup' not in rules_of(iss), 'rules=%s' % rules_of(iss))
    check('注入后问题条数不变（%d → %d）' % (base_count, len(iss)),
          len(iss) == base_count, 'before=%d after=%d' % (base_count, len(iss)))
    print('  [info] 说明：%s；页边距因未提供 spec 标准而整体未判（见 meta.notes）' % info)

    # ---------------------------------------------------- 2 page-numbering
    section('2 注入：删掉第 1 节 w:pgNumType → page-numbering')
    f = fresh(src, 'pagenum')
    info = inject_delete_pgNumType(f, 1)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'page-numbering']
    check('检出 page-numbering', bool(hit), 'rules=%s' % rules_of(iss))
    check('定位到节（location.index 指向第 1 节首个段落）',
          bool(hit) and hit[0]['location']['index'] is not None)
    check('每条都满足 schema', all(not validate_issue(i) for i in hit))
    print('  [info] 注入：%s；命中 %d 条 page-numbering' % (info, len(hit)))

    # ---------------------------------------------------- 3 header
    section('3 注入：某节页眉加一句不一致文字 → header')
    f = fresh(src, 'header')
    info = inject_header_inconsistent(f, 1)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'header']
    check('检出 header', bool(hit), 'rules=%s' % rules_of(iss))
    check('命中"页眉文字不一致"',
          any('文字不一致' in (i.get('actual') or '') for i in hit),
          'ids=%s' % [i['id'] for i in hit])
    check('每条都满足 schema', all(not validate_issue(i) for i in hit))
    print('  [info] 注入：%s；命中 %d 条 header' % (info, len(hit)))

    # ---------------------------------------------------- 4 fonts / style-compliance
    section('4 注入：某正文段设显式 w:sz(9pt) → fonts 或 style-compliance')
    f = fresh(src, 'fonts')
    info = inject_run_sz(f, 18, '正文文本1')
    iss = audit(f)
    hit_rules = rules_of(iss) & {'fonts', 'style-compliance'}
    check('检出 fonts 或 style-compliance', bool(hit_rules), 'rules=%s' % rules_of(iss))
    hit = [i for i in iss if i['rule'] in hit_rules]
    check('确实有"直接格式/字号不一致"类问题',
          any('sz' in (i['actual'] + i['expected']) or '字号' in (i['actual'] + i['expected'])
              for i in hit),
          'actuals=%s' % [i['actual'][:40] for i in hit])
    check('每条都满足 schema', all(not validate_issue(i) for i in hit))
    print('  [info] 注入：%s；命中规则=%s，共 %d 条' % (info, hit_rules, len(hit)))

    # ---------------------------------------------------- 5 toc
    section('5 注入：清空目录域 instrText → toc')
    f = fresh(src, 'toc')
    info = inject_clear_toc(f)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'toc']
    check('检出 toc', bool(hit), 'rules=%s' % rules_of(iss))
    check('为 error 级（域已失效）',
          bool(hit) and hit[0]['severity'] == 'error',
          'sev=%s' % (hit[0]['severity'] if hit else None))
    check('每条都满足 schema', all(not validate_issue(i) for i in hit))
    print('  [info] 注入：%s；命中 %d 条 toc' % (info, len(hit)))

    # ---------------------------------------------------- 收尾
    print('\n' + '=' * 62)
    print('通过 %d 项，失败 %d 项' % (_ok, _fail))
    if _notes:
        print('失败明细：')
        for n in _notes:
            print('  - ' + n)
    print('临时副本留在：%s' % TMP)
    return 0 if _fail == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
