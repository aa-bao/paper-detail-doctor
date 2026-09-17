#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""
tests/selftest_structure_length.py —— structure-length 的**注入式**验收

为什么不能只跑一遍干净文档就收手
    "干净稿上跑出 0 条 error" 不能证明检查器有效——一个永远返回空列表的函数也行。
    所以这里反过来：**人为造出每种缺陷，确认检查器抓得到**；并独立实现一遍字数统计
    做交叉验证（不用被测函数验自己）。

四个场景（都在临时副本上，绝不碰你的原稿）
    0 基线        干净稿 → 只应有 info 级提示，不得有 error；
                  并打印规则分布、条数、length 三个口径数值；
                  用独立实现交叉验证 length 数字一致。
    1 章节编号断号/跳级  → heading-sequence（把 2.3 改成 2.4 造成跳号）
    2 图表编号重复      → figure-table（复制一个表格并改题注编号为重复）
    3 摘要超上限        → abstract-layout（往中文摘要段追加文字使其超 1000 字）

怎么跑
    cd F:\Coding\Project\paper-detail-doctor
    "C:/Users/Tian/.workbuddy/binaries/python/envs/default/Scripts/python.exe" tests/selftest_structure_length.py
"""
import copy
import importlib.util
import os
import re
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
sys.path.insert(0, PKG)

from docx.oxml.ns import qn as _qn  # noqa: E402
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
from shared.lib.docx_ops import (                                    # noqa: E402
    Doc, W as OPS_W, parse_xml, nsdecls,
)
from shared.lib.docx_scan import norm                                # noqa: E402

DEFAULT_SAMPLE = r'F:\Coding\work\0813-城市家庭\初稿.docx'
TMP = tempfile.mkdtemp(prefix='pdd-struct-selftest-')

_ok, _fail, _notes = 0, 0, []


def check(name, cond, detail=''):
    global _ok, _fail
    if cond:
        _ok += 1
        print('  [ok]   %s' % name)
    else:
        _fail += 1
        print('  [FAIL] %s   %s' % (name, detail))
        _notes.append('%s :: %s' % (name, detail))


def section(t):
    print('\n=== %s ===' % t)


# ---------------------------------------------------------------- 加载被审计模块
def load_audit_mod():
    p = os.path.join(PKG, 'skills', 'structure-length', 'scripts', 'audit.py')
    spec = importlib.util.spec_from_file_location('structure_length_audit', p)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


AUD = None


def audit(docx, spec=None):
    return AUD.audit(docx, spec, verbose=False)['issues']


def rules_of(issues):
    return {i['rule'] for i in issues}


def fresh(src, tag):
    p = os.path.join(TMP, tag + '.docx')
    shutil.copy2(src, p)
    return p


# ---------------------------------------------------------------- 独立字数统计（不用被测函数）
# 复制被测模块的区域/块定义，单独实现一遍，用来和被测结果做交叉验证。
def _indep_regions(docx):
    from docx import Document
    paras = list(Document(docx).paragraphs)
    first_h1 = toc = ref = app = None
    for i, p in enumerate(paras, 1):
        t = p.text.strip()
        nt = norm(t)
        if first_h1 is None and p.style.name == '一级标题':
            first_h1 = i
        if toc is None and nt == '目录':
            toc = i
        if ref is None and nt in {'参考文献', 'references', 'bibliography', '主要参考文献'}:
            ref = i
        if app is None and t.startswith('附录'):
            app = i
    return {'first_h1': first_h1, 'toc': toc, 'ref': ref, 'appendix': app}, paras


def _indep_walk(docx):
    from docx import Document
    doc = Document(docx)
    paras = list(doc.paragraphs)
    el_to_idx = {p._element: i for i, p in enumerate(paras, 1)}
    regions, _ = _indep_regions(docx)
    blocks = []
    cur = 'front'
    for child in doc.element.body:
        tag = child.tag
        if tag == W + 'p':
            idx = el_to_idx.get(child)
            txt = ''.join(n.text or '' for n in child.iter(W + 't'))
            if idx is not None:
                if regions['first_h1'] and idx < regions['first_h1']:
                    region = 'front'
                elif regions['ref'] and idx >= regions['ref']:
                    region = 'appendix' if (regions['appendix'] and idx >= regions['appendix']) else 'ref'
                else:
                    region = 'body'
            else:
                region = cur
            blocks.append((txt, region))
            cur = region
        elif tag == W + 'tbl':
            blocks.append(('\n'.join(c.text for row in doc.tables[0].rows
                                      for c in row.cells) if False else
                           ''.join(''.join(t.text or '' for t in tc.iter(W + 't'))
                                   for tc in child.iter(W + 'tc')), cur))
    return blocks


def _indep_count(text):
    hanzi = len(re.findall(r'[\u4e00-\u9fff]', text))
    full = len(re.findall(r'[\u2E80-\u9FFF\u3000-\u303F\uFF00-\uFFEF]', text))
    nows = len(re.sub(r'[\s\u3000]', '', text))
    return hanzi, full, nows


def indep_length(docx):
    blocks = _indep_walk(docx)
    full = _indep_count(''.join(b[0] for b in blocks))
    body = _indep_count(''.join(b[0] for b in blocks if b[1] == 'body'))
    # 分章
    doc = __import__('docx').Document(docx)
    paras = list(doc.paragraphs)
    regions, _ = _indep_regions(docx)
    chapters = []
    cur = None
    order = 0
    cur_region = 'front'
    for child in doc.element.body:
        if child.tag == W + 'p':
            idx = None
            for i, p in enumerate(paras, 1):
                if p._element is child:
                    idx = i
                    break
            txt = ''.join(n.text or '' for n in child.iter(W + 't'))
            if idx is not None:
                if regions['first_h1'] and idx < regions['first_h1']:
                    region = 'front'
                elif regions['ref'] and idx >= regions['ref']:
                    region = 'appendix' if (regions['appendix'] and idx >= regions['appendix']) else 'ref'
                else:
                    region = 'body'
            else:
                region = cur_region
            if region == 'body' and paras[idx - 1].style.name == '一级标题':
                if cur is not None:
                    chapters.append(cur)
                cur = [paras[idx - 1].text.strip(), 0, 0, 0]
            if cur is not None and region == 'body':
                h, f, n = _indep_count(txt)
                cur[1] += h
                cur[2] += f
                cur[3] += n
        elif child.tag == W + 'tbl':
            if cur is not None:
                tbl_txt = ''.join(''.join(t.text or '' for t in tc.iter(W + 't'))
                                  for tc in child.iter(W + 'tc'))
                h, f, n = _indep_count(tbl_txt)
                cur[1] += h
                cur[2] += f
                cur[3] += n
    if cur is not None:
        chapters.append(cur)
    return {
        'with_ref_appendix': {'hanzi': full[0], 'fullwidth': full[1], 'no_whitespace': full[2]},
        'without_ref_appendix': {'hanzi': body[0], 'fullwidth': body[1], 'no_whitespace': body[2]},
        'by_chapter': [{'chapter': c[0], 'hanzi': c[1]} for c in chapters],
    }


# ---------------------------------------------------------------- 注入器（直接改临时副本）

def set_para_text(p_el, new_text):
    """把段落可见文字整体替换为 new_text（保留 rPr； auditing 只读不在乎版式）。"""
    runs = p_el.findall(W + 'r')
    first_t = None
    for r in runs:
        t = r.find(W + 't')
        if t is not None:
            if first_t is None:
                first_t = t
            else:
                r.remove(t)
    if first_t is None:
        r = parse_xml('<w:r %s/>' % nsdecls('w'))
        t = r.makeelement(W + 't', {})
        r.append(t)
        p_el.append(r)
        first_t = t
    first_t.text = new_text
    if new_text != new_text.strip():
        first_t.set(XML_SPACE, 'preserve')


def inject_heading_jump(docx):
    """把 2.3 章的二级标题编号改成 2.4，造成断号（缺 2.3）+ 与 2.3.x 跳级。"""
    d = Doc(docx)
    target = None
    for p in d.paragraphs():
        if p.style.name == '二级标题' and re.match(r'^2\.3\b', p.text.strip()):
            target = p
            break
    if target is None:
        raise RuntimeError('找不到 2.3 二级标题')
    new = re.sub(r'^2\.3\b', '2.4', target.text.strip(), count=1)
    set_para_text(target._element, new)
    d.save_atomic()
    return target.text.strip()


def inject_duplicate_table(docx):
    """复制唯一那个表格，并在它前面放一个编号重复的题注（表 1-1），触发 figure-table。"""
    d = Doc(docx)
    body = d.doc.element.body
    tbl = None
    for child in body:
        if child.tag == W + 'tbl':
            tbl = child
            break
    if tbl is None:
        raise RuntimeError('找不到表格')
    new_tbl = copy.deepcopy(tbl)
    cap = parse_xml('<w:p %s/>' % nsdecls('w'))
    r = cap.makeelement(W + 'r', {})
    t = r.makeelement(W + 't', {})
    t.text = '表 1-1 受访者基本情况（副本）'
    r.append(t)
    cap.append(r)
    at = list(body).index(tbl)
    body.insert(at + 1, cap)
    body.insert(at + 2, new_tbl)
    d.save_atomic()
    return '表 1-1（副本）'


def inject_abstract_overflow(docx):
    """往中文摘要段追加大量汉字，使其超过 default 上限 1000。"""
    d = Doc(docx)
    cn = [i for i, p in enumerate(d.paragraphs(), 1) if p.style.name == '摘要标题'][0]
    target = d.paragraphs()[cn]   # 0-based 索引 cn 即第 cn+1 段（摘要正文）
    add = '本研究进一步发现城市家庭代际情感资源的分配呈现出明显的结构性不对称特征，' * 60
    set_para_text(target._element, target.text + add)
    d.save_atomic()
    return len(re.findall(r'[\u4e00-\u9fff]', target.text + add))


# ---------------------------------------------------------------- 场景

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SAMPLE
    global AUD
    AUD = load_audit_mod()
    if not os.path.exists(src):
        print('样本不存在：%s' % src)
        return 2
    print('样本：%s\n临时目录：%s' % (src, TMP))

    # ---------------------------------------------------- 0 基线
    section('0 基线：干净稿 + length 交叉验证')
    base = fresh(src, 'base')
    res = AUD.audit(base, verbose=False)
    iss = res['issues']
    ls = res['meta']['length']
    print('\n--- 规则分布 ---')
    from collections import Counter
    print('  总条数 = %d' % len(iss))
    for rule, n in sorted(Counter(i['rule'] for i in iss).items()):
        print('    %-16s %d' % (rule, n))
    print('  severity:', dict(Counter(i['severity'] for i in iss)))
    print('  standard_source:', dict(Counter(i['standard_source'] for i in iss)))
    print('\n--- length 三个口径 ---')
    print('  含参考文献与附录 :', ls['with_ref_appendix'])
    print('  不含（仅正文）   :', ls['without_ref_appendix'])
    print('  按章:')
    for c in ls['by_chapter']:
        print('    %-42s hanzi=%d (%.1f%%)' % (c['chapter'], c['hanzi'], c['pct']))

    check('基线无 error', not [i for i in iss if i['severity'] == 'error'],
          'errors=%s' % [(i['rule'], i['actual'][:40]) for i in iss if i['severity'] == 'error'])
    check('基线 heading-sequence 不报错', 'heading-sequence' not in rules_of(iss))
    check('基线 figure-table 不报错', 'figure-table' not in rules_of(iss))
    check('基线 length 有统计', ls is not None and ls['without_ref_appendix']['hanzi'] > 0)
    check('分章汉字数之和 == 正文汉字数',
          sum(c['hanzi'] for c in ls['by_chapter']) == ls['without_ref_appendix']['hanzi'],
          '%d vs %d' % (sum(c['hanzi'] for c in ls['by_chapter']),
                        ls['without_ref_appendix']['hanzi']))

    # 独立交叉验证
    ind = indep_length(base)
    check('交叉验证·全文档汉字数一致',
          ind['with_ref_appendix']['hanzi'] == ls['with_ref_appendix']['hanzi'],
          '%d vs %d' % (ind['with_ref_appendix']['hanzi'], ls['with_ref_appendix']['hanzi']))
    check('交叉验证·全文档全角字符数一致',
          ind['with_ref_appendix']['fullwidth'] == ls['with_ref_appendix']['fullwidth'],
          '%d vs %d' % (ind['with_ref_appendix']['fullwidth'], ls['with_ref_appendix']['fullwidth']))
    check('交叉验证·全文档不含空白字符数一致',
          ind['with_ref_appendix']['no_whitespace'] == ls['with_ref_appendix']['no_whitespace'],
          '%d vs %d' % (ind['with_ref_appendix']['no_whitespace'], ls['with_ref_appendix']['no_whitespace']))
    check('交叉验证·正文汉字数一致',
          ind['without_ref_appendix']['hanzi'] == ls['without_ref_appendix']['hanzi'],
          '%d vs %d' % (ind['without_ref_appendix']['hanzi'], ls['without_ref_appendix']['hanzi']))
    check('交叉验证·分章汉字数一致',
          [c['hanzi'] for c in ind['by_chapter']] == [c['hanzi'] for c in ls['by_chapter']],
          '%s vs %s' % ([c['hanzi'] for c in ind['by_chapter']],
                        [c['hanzi'] for c in ls['by_chapter']]))
    # 量级校验：正文约 37461 字；口径不同会有差异，但必须在合理量级
    body_hanzi = ls['without_ref_appendix']['hanzi']
    check('量级校验·正文字数在合理量级（2万~4.5万）',
          20000 <= body_hanzi <= 45000, 'body_hanzi=%d' % body_hanzi)

    # 只读性
    from shared.lib.ooxml_guard import sha256_file
    h0 = sha256_file(base)
    AUD.audit(base, verbose=False)
    check('审计是只读的（哈希不变）', sha256_file(base) == h0)

    # ---------------------------------------------------- 1 章节编号
    section('1 注入：章节编号断号/跳级（2.3 → 2.4）→ heading-sequence')
    f = fresh(src, 'head')
    old = inject_heading_jump(f)
    print('  注入：%s' % old)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'heading-sequence']
    check('检出 heading-sequence', bool(hit), 'rules=%s' % rules_of(iss))
    check('有 error 级（断号）',
          any(i['severity'] == 'error' for i in hit),
          'sev=%s' % [i['severity'] for i in hit][:6])
    check('描述写明"断号/缺失"',
          any('断号' in i['actual'] or '缺失' in i['actual'] for i in hit))
    check('action 为 None（全局手术不自动改）',
          all(i['action'] is None for i in hit))

    # ---------------------------------------------------- 2 图表编号
    section('2 注入：复制表格并改题注编号为重复 → figure-table')
    f = fresh(src, 'fig')
    cap = inject_duplicate_table(f)
    print('  注入题注：%s' % cap)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'figure-table']
    check('检出 figure-table', bool(hit), 'rules=%s' % rules_of(iss))
    check('报出编号重复', any('重复' in i['actual'] for i in hit),
          'actuals=%s' % [i['actual'][:40] for i in hit])
    check('action 为 None', all(i['action'] is None for i in hit))

    # ---------------------------------------------------- 3 摘要超上限
    section('3 注入：中文摘要追加文字超上限 → abstract-layout')
    f = fresh(src, 'abs')
    new_cn = inject_abstract_overflow(f)
    print('  注入后中文摘要汉字数 ≈ %d' % new_cn)
    iss = audit(f)
    hit = [i for i in iss if i['rule'] == 'abstract-layout']
    check('检出 abstract-layout', bool(hit), 'rules=%s' % rules_of(iss))
    over = [i for i in hit if '超出上限' in i['actual']]
    check('报出"中文摘要超出上限"', bool(over),
          'actuals=%s' % [i['actual'][:40] for i in hit])
    check('超上限为 error 级', bool(over) and over[0]['severity'] == 'error',
          'sev=%s' % (over[0]['severity'] if over else None))
    check('上限来源标注 default（未把范文实测字数当阈值）',
          bool(over) and over[0]['standard_source'] == 'default',
          'src=%s' % (over[0]['standard_source'] if over else None))
    check('action 为 None', all(i['action'] is None for i in hit))

    # ---------------------------------------------------- 收尾
    print('\n' + '=' * 62)
    print('通过 %d 项，失败 %d 项' % (_ok, _fail))
    if _notes:
        print('失败明细：')
        for n in _notes:
            print('  - ' + n)
    print('临时副本留在：%s' % TMP)
    return 0 if _fail == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
